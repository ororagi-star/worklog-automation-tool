from copy import deepcopy
from datetime import date, datetime, time
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from automate_worklog import fill_worklog_dates
from config_store import ConfigConflict, ConfigStore, SessionConfigStore
from worklog_config import check_source_ownership, default_config, validate_config


BASE = Path(__file__).resolve().parent


class SettingsTests(unittest.TestCase):
    def setUp(self):
        self.config = default_config(BASE)

    def test_persistence_conflict_and_restore(self):
        with TemporaryDirectory(dir=BASE / "test") as folder:
            path = Path(folder) / "settings.sqlite3"
            store = ConfigStore(path)
            store.initialize(self.config)
            loaded, revision, _ = store.load()
            old_id = loaded["templates"][0]["items"][0]["id"]
            loaded["templates"][0]["items"][0]["name"] = "변경된 항목"
            store.save(loaded, revision)
            with self.assertRaises(ConfigConflict):
                store.save(self.config, revision)
            reopened = ConfigStore(path)
            reopened.initialize(self.config)
            current, revision, _ = reopened.load()
            self.assertEqual(revision, 2)
            self.assertEqual(current["templates"][0]["items"][0]["id"], old_id)
            self.assertEqual(current["templates"][0]["items"][0]["name"], "변경된 항목")
            reopened.restore(1, revision)
            restored, revision, _ = reopened.load()
            self.assertEqual(revision, 3)
            self.assertEqual(restored, self.config)
            self.assertEqual(len(reopened.history()), 3)

    def test_public_preview_is_isolated_and_can_restore(self):
        first, second = SessionConfigStore({}), SessionConfigStore({})
        first.initialize(self.config)
        second.initialize(self.config)
        changed, revision, _ = first.load()
        changed["templates"][0]["items"][0]["name"] = "임시 변경"
        first.save(changed, revision)
        self.assertEqual(second.load()[0], self.config)
        first.restore(1, 2)
        self.assertEqual(first.load()[0], self.config)
        self.assertEqual(first.load()[1], 3)

    def test_validation_prevents_double_counting(self):
        config = deepcopy(self.config)
        for item in config["templates"][0]["items"][:2]:
            item.update(mode="직접 연결", sources=["같은 수업"])
        with self.assertRaises(ValueError):
            validate_config(config)
        config["templates"][0]["items"][0].update(mode="자동 연결", name="같은 수업", match_name="같은 수업", sources=[])
        with self.assertRaises(ValueError):
            check_source_ownership(config, {"같은 수업"})

    def test_backup_validation(self):
        config = deepcopy(self.config)
        config["templates"][0]["items"][0]["target_people"] = -1
        with self.assertRaises(ValueError):
            validate_config(config)
        config = deepcopy(self.config)
        config["templates"][0]["items"][0]["name"] = "=1+1"
        with self.assertRaises(ValueError):
            validate_config(config)

    def test_dynamic_rows_renames_categories_and_multiple_dates(self):
        items = [
            {**self.config["templates"][0]["items"][0], "id": "new-1", "name": "새 이름", "category": "새 분류", "order": 1,
             "mode": "직접 연결", "sources": ["원본1", "원본2"], "target_sessions": 10},
            {**self.config["templates"][0]["items"][1], "id": "new-2", "name": "행사", "category": "행사", "order": 2, "mode": "집계 제외"},
            {**self.config["templates"][0]["items"][2], "id": "new-3", "name": "추가 수업", "category": "새 분류", "order": 3,
             "mode": "직접 연결", "sources": ["원본3"]},
        ]
        with TemporaryDirectory(dir=BASE / "test") as folder:
            root = Path(folder)
            source = Workbook()
            sheet = source.active
            for col, day in [(10, 1), (11, 2)]:
                sheet.cell(5, col, datetime(2026, 9, day))
            for row, name in [(7, "원본1"), (8, "원본2"), (9, "원본3")]:
                sheet.cell(row, 2, name)
                sheet.cell(row, 10, time(9))
            sheet.cell(7, 11, time(9))
            source.save(root / "source.xlsx")
            summary = fill_worklog_dates(root / "source.xlsx", BASE / "worklog_set1.xlsx", root / "out.xlsx",
                                         [date(2026, 9, 1), date(2026, 9, 2)], config_items=items)
            self.assertEqual(summary["date_summaries"][0]["matched_rows"][0]["daily_sessions"], 1)
            self.assertEqual(summary["date_summaries"][0]["matched_rows"][0]["daily_people"], 2)
            self.assertEqual(summary["date_summaries"][1]["matched_rows"][0]["month_sessions"], 2)
            workbook = load_workbook(root / "out.xlsx")
            sheet = workbook.active
            self.assertEqual(sheet["A7"].value, "새 분류")
            self.assertEqual(sheet["B7"].value, "새 이름")
            self.assertEqual(sheet["B8"].value, "추가 수업")
            self.assertEqual(sheet["B9"].value, "행사")
            self.assertIn("A7:A8", [str(area) for area in sheet.merged_cells.ranges])
            self.assertIsNone(sheet["E9"].value)
            self.assertEqual(sheet["E6"].value, 2)
            self.assertEqual(sheet["F6"].value, 3)
            self.assertEqual(sheet["C6"].value, "=SUM(C7:C9)")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
