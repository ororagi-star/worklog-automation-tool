import unittest
from datetime import date

from openpyxl import Workbook

from automate_worklog import fill_sheet_block, match_programs, normalize_key


class ProgramMatchingTests(unittest.TestCase):
    def test_current_names_override_old_aliases(self):
        for item in ["일본어 첫걸음", "중국어 첫걸음", "한글교실 중급", "생활 영어"]:
            with self.subTest(item=item):
                self.assertEqual(match_programs(item, {item}), [item])
        self.assertEqual(
            match_programs("일본어 첫걸음", {"일본어 첫걸음", "일본어(초급)"}),
            ["일본어 첫걸음"],
        )

    def test_known_renames_work_in_both_directions(self):
        self.assertEqual(match_programs("일본어 첫걸음", {"일본어(초급)"}), ["일본어(초급)"])
        self.assertEqual(match_programs("일본어(초급)", {"일본어 첫걸음"}), ["일본어 첫걸음"])
        self.assertEqual(match_programs("영어 첫걸금", {"영어 첫걸음"}), ["영어 첫걸음"])

    def test_spacing_and_class_number_variants(self):
        self.assertEqual(match_programs(" 일본어첫걸음 ", {"일본어 첫걸음"}), ["일본어 첫걸음"])
        self.assertEqual(match_programs("스마트폰(기초1)", {"스마트폰(기초1반)"}), ["스마트폰(기초1반)"])
        self.assertNotEqual(normalize_key("일반요가"), normalize_key("일요가"))

    def test_levels_and_substrings_are_not_guessed(self):
        self.assertEqual(match_programs("한글교실(중급)", {"한글교실(초급1반)", "한글교실(초급2반)"}), [])
        self.assertEqual(match_programs("시니어 디지털존", {"디지털존"}), [])
        self.assertEqual(match_programs("새 수업", {"새 수업 심화"}), [])

    def test_ambiguous_aliases_are_not_combined(self):
        self.assertEqual(match_programs("영어 첫걸금", {"영어 첫걸음", "영어(초급)"}), [])

    def test_unmatched_row_does_not_keep_template_numbers(self):
        sheet = Workbook().active
        sheet.cell(7, 2, "새 수업")
        for col in range(5, 9):
            sheet.cell(7, col, 99)
        summary = fill_sheet_block(sheet, date(2026, 9, 4), 1, 7, {}, {}, set())
        self.assertEqual([sheet.cell(7, col).value for col in range(5, 9)], [None] * 4)
        self.assertEqual(summary["unmatched_rows"][0]["item"], "새 수업")


if __name__ == "__main__":
    unittest.main()
