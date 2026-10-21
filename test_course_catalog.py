from copy import deepcopy
from datetime import date, datetime, time
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from automate_worklog import fill_worklog_dates
from config_store import ConfigStore
from course_catalog import add_category, assign_courses, delete_category, delete_course, place_course, remove_course_everywhere, rename_category, save_course
from worklog_config import default_config, resolve_sources, validate_config


BASE = Path(__file__).resolve().parent


class CatalogueTests(unittest.TestCase):
    def setUp(self):
        self.config = default_config(BASE)

    def test_register_before_assignment_and_persist_empty_category(self):
        config = add_category(self.config, "신규 활동")
        original_rows = deepcopy(config["templates"])
        config = save_course(config, {"category": "신규 활동", "name": "새 수업"})
        self.assertEqual(config["templates"], original_rows)
        course = config["courses"][-1]
        self.assertEqual(course["name"], "새 수업")
        with TemporaryDirectory(dir=BASE / "test") as folder:
            store = ConfigStore(Path(folder) / "settings.sqlite3")
            store.initialize(config)
            store.save(add_category(config, "아직 수업 없는 분류"), 1)
            loaded, _, _ = store.load()
            self.assertIn("아직 수업 없는 분류", loaded["categories"])
            self.assertEqual(loaded["courses"][-1]["id"], course["id"])
        config = delete_course(config, course["id"])
        config = delete_category(config, "신규 활동")
        self.assertNotIn("신규 활동", config["categories"])

    def test_assignment_rename_and_output_from_catalogue(self):
        config = add_category(self.config, "새 카테고리")
        config = save_course(config, {"category": "새 카테고리", "name": "새 수업", "mode": "직접 연결", "sources": ["원본 수업"]})
        course = config["courses"][-1]
        config = assign_courses(config, "worklog_set1", [course["id"]])
        item_id = config["templates"][0]["items"][0]["id"]
        config = save_course(config, {**course, "name": "바뀐 수업명"})
        config = rename_category(config, "새 카테고리", "바뀐 카테고리")
        item = config["templates"][0]["items"][0]
        self.assertEqual(item["id"], item_id)
        self.assertEqual(item["sources"], ["원본 수업"])
        with self.assertRaises(ValueError):
            delete_course(config, course["id"])
        with self.assertRaises(ValueError):
            delete_category(config, "바뀐 카테고리")
        with TemporaryDirectory(dir=BASE / "test") as folder:
            root = Path(folder)
            source = Workbook()
            source.active.cell(5, 10, datetime(2026, 9, 4))
            source.active.cell(7, 2, "원본 수업")
            source.active.cell(7, 10, time(9))
            source.save(root / "source.xlsx")
            fill_worklog_dates(root / "source.xlsx", BASE / "worklog_set1.xlsx", root / "out.xlsx",
                               [date(2026, 9, 4)], config_items=config["templates"][0]["items"])
            workbook = load_workbook(root / "out.xlsx")
            self.assertEqual(workbook.active["A7"].value, "바뀐 카테고리")
            self.assertEqual(workbook.active["B7"].value, "바뀐 수업명")
            self.assertEqual(workbook.active["F7"].value, 1)
            workbook.close()

    def test_old_backup_migrates_without_changing_row_ids(self):
        old = deepcopy(self.config)
        old.pop("categories")
        old.pop("courses")
        for template in old["templates"]:
            for item in template["items"]:
                item.pop("course_id")
        upgraded = validate_config(old)
        self.assertEqual(len(upgraded["courses"]), 90)
        self.assertEqual(upgraded["templates"][0]["items"][0]["id"], "worklog_set1-7")
        self.assertEqual(upgraded["templates"][0]["items"][0]["course_id"], "worklog_set1-7")

    def test_automatic_connection_survives_display_name_change(self):
        config = save_course(self.config, {"category": "언어활동", "name": "처음 원본명"})
        course = config["courses"][-1]
        config = assign_courses(config, "worklog_set1", [course["id"]])
        config = save_course(config, {**course, "name": "새 표시명"})
        item = config["templates"][0]["items"][0]
        self.assertEqual(item["name"], "새 표시명")
        self.assertEqual(resolve_sources(item, {"처음 원본명"}), ["처음 원본명"])

    def test_add_move_and_remove_in_one_action(self):
        config = place_course(self.config, {"category": "새 활동", "name": "새 수업"}, ["worklog_set1"])
        course = config["courses"][-1]
        self.assertIn("새 활동", config["categories"])
        self.assertEqual(config["templates"][0]["items"][-1]["course_id"], course["id"])
        config = place_course(config, {**course, "name": "수정한 이름"}, ["worklog_set2"])
        self.assertFalse(any(i["course_id"] == course["id"] for i in config["templates"][0]["items"]))
        self.assertEqual(config["templates"][1]["items"][-1]["name"], "수정한 이름")
        config = remove_course_everywhere(config, course["id"])
        self.assertFalse(any(i["course_id"] == course["id"] for t in config["templates"] for i in t["items"]))
        self.assertFalse(any(c["id"] == course["id"] for c in config["courses"]))


if __name__ == "__main__":
    unittest.main()
