from copy import copy
from math import ceil

from openpyxl.styles import Alignment


def copy_region(source, target, first, last, destination):
    offset = destination - first
    for row in range(first, last + 1):
        target.row_dimensions[row + offset].height = source.row_dimensions[row].height
        for col in range(1, 13):
            old = source.cell(row, col)
            new = target.cell(row + offset, col)
            new.value = old.value
            new._style = copy(old._style)
    for area in source.merged_cells.ranges:
        if first <= area.min_row and area.max_row <= last:
            target.merge_cells(start_row=area.min_row + offset, end_row=area.max_row + offset,
                               start_column=area.min_col, end_column=area.max_col)


def build_configured_sheet(workbook, source, items):
    sheet = workbook.create_sheet("configured")
    copy_region(source, sheet, 1, 6, 1)
    for key, dimension in source.column_dimensions.items():
        sheet.column_dimensions[key] = copy(dimension)
    sheet.sheet_format = copy(source.sheet_format)
    sheet.page_setup = copy(source.page_setup)
    sheet.page_margins = copy(source.page_margins)
    sheet.print_options = copy(source.print_options)
    sheet.sheet_properties = copy(source.sheet_properties)
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    active = sorted((item for item in items if item["enabled"]), key=lambda item: item["order"])
    category_order = list(dict.fromkeys(item["category"] for item in active))
    active.sort(key=lambda item: (category_order.index(item["category"]), item["order"]))
    row_items = {}
    group_start = 7
    previous = None
    for row, item in enumerate(active, 7):
        for col in range(1, 13):
            sheet.cell(row, col)._style = copy(source.cell(10, col)._style)
        sheet.row_dimensions[row].height = max(24, 16 * ceil(len(item["name"]) / 13))
        sheet.cell(row, 2, item["name"])
        sheet.cell(row, 2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.cell(row, 3, item["target_sessions"])
        sheet.cell(row, 4, item["target_people"])
        sheet.merge_cells(start_row=row, start_column=11, end_row=row, end_column=12)
        if previous != item["category"]:
            if previous is not None and row - group_start > 1:
                sheet.merge_cells(start_row=group_start, end_row=row - 1, start_column=1, end_column=1)
            group_start = row
            sheet.cell(row, 1, item["category"])
            sheet.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            previous = item["category"]
        row_items[row] = item
    last_item_row = 6 + len(active)
    if last_item_row > group_start:
        sheet.merge_cells(start_row=group_start, end_row=last_item_row, start_column=1, end_column=1)
    for col in (3, 4):
        letter = sheet.cell(6, col).column_letter
        sheet.cell(6, col, f"=SUM({letter}7:{letter}{last_item_row})")
    footer_start = next((row for row in range(7, source.max_row + 1)
                         if str(source.cell(row, 1).value or "").strip() == "업무내용"), None)
    if footer_start:
        copy_region(source, sheet, footer_start, source.max_row, last_item_row + 1)
    sheet.print_area = f"A1:L{sheet.max_row}"
    sheet.print_title_rows = "4:5"
    title = source.title
    index = workbook.index(source)
    workbook.remove(source)
    sheet.title = title
    workbook.move_sheet(sheet, offset=index - workbook.index(sheet))
    return sheet, row_items
