from __future__ import annotations

from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from automate_worklog import SKIP_WORKLOG_ITEMS, normalize_text


TEMPLATES = [("worklog_set1", "평생교육사업1"), ("worklog_set2", "평생교육사업2")]
MODES = ["자동 연결", "직접 연결", "집계 제외"]
SESSION_RULES = ["하루 1건", "수업별 합산"]


def default_config(base_dir: Path) -> dict:
    templates = []
    for key, label in TEMPLATES:
        workbook = load_workbook(base_dir / f"{key}.xlsx")
        sheet = workbook.active
        items = []
        for row in range(7, sheet.max_row + 1):
            name = normalize_text(sheet.cell(row, 2).value)
            if not name:
                continue
            category = sheet.cell(row, 1).value
            for area in sheet.merged_cells.ranges:
                if area.min_col == 1 and area.min_row <= row <= area.max_row:
                    category = sheet.cell(area.min_row, 1).value
                    break
            items.append({
                "id": f"{key}-{row}", "category": normalize_text(category).replace(" ", ""),
                "name": name, "order": len(items) + 1, "enabled": True,
                "mode": "집계 제외" if name in SKIP_WORKLOG_ITEMS else "자동 연결",
                "sources": [], "session_rule": "하루 1건",
                "target_sessions": sheet.cell(row, 3).value or 0,
                "target_people": sheet.cell(row, 4).value or 0,
            })
        templates.append({"key": key, "label": label, "items": items})
        workbook.close()
    return validate_config({"schema_version": 1, "templates": templates})


def clean_item(source: dict) -> dict:
    if not isinstance(source, dict):
        raise ValueError("수업 설정을 확인해 주세요.")
    item = {"id": source.get("id") or str(uuid4())}
    if not isinstance(item["id"], str):
        raise ValueError("항목 ID를 확인해 주세요.")
    for field in ("name", "category"):
        value = normalize_text(source.get(field))
        if not value or value.startswith(("=", "+", "-", "@")):
            raise ValueError("카테고리와 수업명을 입력해 주세요. 이름은 수식 기호로 시작할 수 없습니다.")
        item[field] = value
    item["course_id"] = source.get("course_id") or item["id"]
    item["match_name"] = normalize_text(source.get("match_name")) or item["name"]
    item["enabled"] = bool(source.get("enabled", True))
    item["mode"] = source.get("mode", "자동 연결")
    item["session_rule"] = source.get("session_rule", "하루 1건")
    if item["mode"] not in MODES or item["session_rule"] not in SESSION_RULES:
        raise ValueError(f"집계 설정을 확인해 주세요: {item['name']}")
    sources = source.get("sources") or []
    if not isinstance(sources, list) or any(not isinstance(s, str) for s in sources):
        raise ValueError("원본 수업 연결을 확인해 주세요.")
    item["sources"] = list(dict.fromkeys(normalize_text(s) for s in sources if normalize_text(s)))
    if item["mode"] == "직접 연결" and item["enabled"] and not item["sources"]:
        raise ValueError(f"직접 연결할 원본 수업을 선택해 주세요: {item['name']}")
    for field in ("order", "target_sessions", "target_people"):
        value = source.get(field, 0)
        try:
            number = int(value)
            if isinstance(value, bool) or float(value) != number or number < 0:
                raise ValueError()
        except (ValueError, TypeError, OverflowError):
            raise ValueError("순서와 목표는 0 이상의 정수로 입력해 주세요.") from None
        item[field] = number
    return item


def validate_config(config: dict) -> dict:
    if not isinstance(config, dict) or config.get("schema_version") != 1:
        raise ValueError("지원하지 않는 설정 형식입니다.")
    templates = config.get("templates")
    if not isinstance(templates, list) or len(templates) != len(TEMPLATES):
        raise ValueError("양식 1과 양식 2의 설정이 모두 필요합니다.")
    result = {"schema_version": 1, "templates": []}
    ids, explicit_sources, seen_keys = set(), set(), set()
    for template in templates:
        key = template.get("key")
        if key not in dict(TEMPLATES) or key in seen_keys:
            raise ValueError("양식 구분이 잘못되었거나 중복되었습니다.")
        seen_keys.add(key)
        items = [clean_item(source) for source in template.get("items", [])]
        names = set()
        for item in items:
            if item["id"] in ids or item["name"] in names:
                raise ValueError(f"양식에 ID 또는 수업명이 중복되었습니다: {item['name']}")
            ids.add(item["id"])
            names.add(item["name"])
            if item["mode"] == "직접 연결" and item["enabled"]:
                for program in item["sources"]:
                    if program in explicit_sources:
                        raise ValueError(f"원본 수업을 여러 항목에 중복 연결할 수 없습니다: {program}")
                    explicit_sources.add(program)
        if not any(item["enabled"] for item in items):
            raise ValueError("각 업무일지에는 수업을 하나 이상 선택해 주세요.")
        result["templates"].append({"key": key, "label": dict(TEMPLATES)[key], "items": items})

    # Older backups contain only template rows; promote them to a shared catalogue.
    raw_courses = config.get("courses", [item for t in result["templates"] for item in t["items"]])
    if not isinstance(raw_courses, list):
        raise ValueError("수업 목록을 확인해 주세요.")
    result["courses"] = [clean_item(course) for course in raw_courses]
    course_ids = [course["id"] for course in result["courses"]]
    if len(course_ids) != len(set(course_ids)):
        raise ValueError("수업 목록에 ID가 중복되었습니다.")
    course_names = [(course["category"], course["name"]) for course in result["courses"]]
    if len(course_names) != len(set(course_names)):
        raise ValueError("같은 카테고리에 같은 수업명이 이미 있습니다.")
    categories = config.get("categories", list(dict.fromkeys(course["category"] for course in result["courses"])))
    if not isinstance(categories, list) or any(not isinstance(name, str) for name in categories):
        raise ValueError("카테고리 목록을 확인해 주세요.")
    result["categories"] = [normalize_text(name) for name in categories]
    if any(not name or name.startswith(("=", "+", "-", "@")) for name in result["categories"]):
        raise ValueError("카테고리 이름을 확인해 주세요.")
    if len(set(result["categories"])) != len(result["categories"]):
        raise ValueError("같은 카테고리가 이미 있습니다.")
    catalogue = {course["id"]: course for course in result["courses"]}
    for course in result["courses"]:
        if course["category"] not in result["categories"]:
            raise ValueError(f"등록되지 않은 카테고리입니다: {course['category']}")
    for template in result["templates"]:
        for item in template["items"]:
            if not isinstance(item["course_id"], str) or item["course_id"] not in catalogue:
                raise ValueError(f"수업 목록에 없는 항목입니다: {item['name']}")
    return result


def resolve_sources(item: dict, available: set[str]) -> list[str]:
    from automate_worklog import match_programs

    if item["mode"] == "집계 제외":
        return []
    if item["mode"] == "직접 연결":
        return [name for name in item["sources"] if name in available]
    return match_programs(item.get("match_name", item["name"]), available)


def check_source_ownership(config: dict, available: set[str]) -> None:
    owners = {}
    for template in config["templates"]:
        for item in template["items"]:
            if not item["enabled"]:
                continue
            for name in resolve_sources(item, available):
                if name in owners:
                    raise ValueError(f"원본 수업 '{name}'이 '{owners[name]}'와 '{item['name']}'에 중복 연결됩니다.")
                owners[name] = item["name"]
